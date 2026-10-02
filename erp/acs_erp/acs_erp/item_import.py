"""Plan 017: import the item master from the old system's Excel export.

Run by hand only (never from after_install / after_migrate):

    bench --site frontend execute acs_erp.item_import.run --kwargs "{'path': '/tmp/items.xlsx', 'apply': False}"

Dry run by default; apply=True writes and commits once. Re-runs update only the mapped
fields that differ and never delete anything.
"""
import re

PANEL_ITEMS = {"PLC", "SCADA", "HMI"}
SIEMENS_PREFIXES = ("6ES7", "6AV", "6EP", "6SL", "3RT", "3RV")
GROUP_COLUMNS = ("CatDesc", "Group", "GroupDesc")
# Columns the mapping reads. If a new export lacks one, stop and look (plan 017: don't guess).
REQUIRED_COLUMNS = ("ItemCode", "ItemDesc", "DetailDesc", "PartNo", "Active", "MakeDesc") + GROUP_COLUMNS
NAME_MAX = 140


def _text(value):
    """Trimmed text with runs of whitespace collapsed to one space; None for empty."""
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def _active(value):
    if value is None or value == "":
        return True
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("1", "true", "yes", "y")


def clean_row(header, row, row_number):
    """One spreadsheet row → the Item fields, or {"error": ...} / {"skip": ...} naming the row."""
    columns = {(_text(h) or ""): i for i, h in enumerate(header)}

    def cell(name):
        i = columns.get(name)
        return row[i] if i is not None and i < len(row) else None

    code = _text(cell("ItemCode"))
    if not code:
        return {"error": f"Row {row_number}: ItemCode is empty."}
    name = _text(cell("ItemDesc"))
    if not name:
        return {"error": f"Row {row_number}: ItemDesc is empty."}
    if code in PANEL_ITEMS:
        return {"skip": f"Row {row_number}: {code} is a panel item used by projects; not imported."}

    part_no = _text(cell("PartNo"))
    brand = _text(cell("MakeDesc"))
    if not brand and part_no and part_no.upper().startswith(SIEMENS_PREFIXES):
        brand = "Siemens"

    item_name = name[:NAME_MAX]
    return {
        "item_code": code,
        "item_name": item_name,
        "description": _text(cell("DetailDesc")) or item_name,
        "custom_maker_part_no": part_no,
        "brand": brand,
        "item_group_path": [g for g in (_text(cell(c)) for c in GROUP_COLUMNS) if g],
        "disabled": 0 if _active(cell("Active")) else 1,
        "stock_uom": "Nos",
        "is_stock_item": 1,
    }


def clean_rows(header, rows):
    """All data rows (spreadsheet row 2 onward) → (items, problems). A repeated ItemCode is reported once more."""
    items, problems, first_seen = [], [], {}
    for offset, row in enumerate(rows):
        row_number = offset + 2
        if not any(v not in (None, "") for v in row):
            continue
        result = clean_row(header, row, row_number)
        if "error" in result or "skip" in result:
            problems.append(result)
            continue
        code = result["item_code"]
        if code in first_seen:
            problems.append({"error": f"Row {row_number}: ItemCode {code} is repeated (first on row {first_seen[code]}); not imported."})
            continue
        first_seen[code] = row_number
        items.append(result)
    return items, problems


UPDATE_FIELDS = ("item_name", "description", "custom_maker_part_no", "brand", "item_group", "disabled")


def _same(a, b):
    return (a if a not in ("", None) else None) == (b if b not in ("", None) else None)


def run(path, apply=False):
    """Reads the first sheet of `path` and loads it as Items. Returns the report dict (also printed)."""
    import frappe
    from openpyxl import load_workbook

    report = {"created": 0, "updated": 0, "unchanged": 0, "skipped": 0, "errors": 0, "new_groups": [], "new_brands": [], "lines": []}
    rows = list(load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    header, data = (rows[0], rows[1:]) if rows else ((), [])
    present = {(_text(h) or "") for h in header}
    missing = [c for c in REQUIRED_COLUMNS if c not in present]
    if missing:
        report["errors"] = 1
        report["lines"].append(f"The file has no column {', '.join(missing)}. Nothing was imported; check the export.")
        return _finish(report, apply)

    items, problems = clean_rows(header, data)
    for problem in problems:
        report["skipped" if "skip" in problem else "errors"] += 1
        report["lines"].append(problem.get("skip") or problem["error"])

    try:
        # Item groups parent-first under "All Item Groups"; the last name of a path is a leaf.
        for item in items:
            parent = "All Item Groups"
            for depth, name in enumerate(item["item_group_path"]):
                if not frappe.db.exists("Item Group", name) and name not in report["new_groups"]:
                    report["new_groups"].append(name)
                    if apply:
                        frappe.get_doc({
                            "doctype": "Item Group",
                            "item_group_name": name,
                            "parent_item_group": parent,
                            "is_group": 1 if depth < len(item["item_group_path"]) - 1 else 0,
                        }).insert()
                parent = name
        for brand in dict.fromkeys(i["brand"] for i in items if i["brand"]):
            if not frappe.db.exists("Brand", brand):
                report["new_brands"].append(brand)
                if apply:
                    frappe.get_doc({"doctype": "Brand", "brand": brand}).insert()

        verb = "" if apply else "would "
        for item in items:
            fields = {k: item[k] for k in UPDATE_FIELDS if k != "item_group"}
            fields["item_group"] = item["item_group_path"][-1] if item["item_group_path"] else "All Item Groups"
            code = item["item_code"]
            if frappe.db.exists("Item", code):
                doc = frappe.get_doc("Item", code)
                changed = {k: v for k, v in fields.items() if not _same(doc.get(k), v)}
                if not changed:
                    report["unchanged"] += 1
                    continue
                report["updated"] += 1
                report["lines"].append(f"{code}: {verb}update {', '.join(changed)}")
                if apply:
                    doc.update(changed)
                    doc.save()
            else:
                report["created"] += 1
                report["lines"].append(f"{code}: {verb}create")
                if apply:
                    frappe.get_doc({
                        "doctype": "Item",
                        "item_code": code,
                        **fields,
                        "stock_uom": item["stock_uom"],
                        "is_stock_item": item["is_stock_item"],
                    }).insert()
        if apply:
            frappe.db.commit()
    except Exception:
        if apply:
            frappe.db.rollback()
        raise
    return _finish(report, apply)


def _finish(report, apply):
    for line in report["lines"]:
        print(line)
    print(
        f"{'Applied' if apply else 'Dry run (nothing written; pass apply=True to write)'}: "
        f"{report['created']} created, {report['updated']} updated, {report['unchanged']} unchanged, "
        f"{report['skipped']} skipped, {report['errors']} errors; "
        f"new item groups: {', '.join(report['new_groups']) or 'none'}; new brands: {', '.join(report['new_brands']) or 'none'}"
    )
    return report
